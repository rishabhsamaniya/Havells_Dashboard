import io
from datetime import datetime
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.analytics.services.analytics_service import AnalyticsService

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name="Segoe UI", size=9, italic=True, color="64748B")
DATA_FONT = Font(name="Segoe UI", size=10, color="1E293B")
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
TOTAL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

CURRENCY_FORMAT = '[$₹-en-IN]#,##,##0.00'
PERCENT_FORMAT = '0.00%'
NUMBER_FORMAT = '#,##,##0'

def apply_auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.number_format in (CURRENCY_FORMAT, PERCENT_FORMAT):
                val += '       '
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

class ExportService:
    @staticmethod
    def generate_consolidated_mis_workbook(filters=None):
        wb = openpyxl.Workbook()
        # Default active sheet
        ws_overview = wb.active
        ws_overview.title = "Executive Overview"

        qs = AnalyticsService.get_filtered_queryset(filters)
        kpis = AnalyticsService.get_executive_kpis(qs)
        now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

        # -------------------------------------------------------------
        # Tab 1: Executive Overview
        # -------------------------------------------------------------
        ws_overview.append(["HAVELLS BILLING INTELLIGENCE & MIS EXECUTIVE SUMMARY"])
        ws_overview.cell(row=1, column=1).font = TITLE_FONT
        ws_overview.append([f"Report Generated: {now_str} | Active Population: {kpis['total_records']} Transactions"])
        ws_overview.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_overview.append([])

        ws_overview.append(["Key Performance Metric", "Value", "Metric Description / Accounting Note"])
        for col_idx in range(1, 4):
            cell = ws_overview.cell(row=4, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center" if col_idx == 2 else "left")

        overview_metrics = [
            ("Total Billing Amount", float(kpis['total_billing']), CURRENCY_FORMAT, "Gross invoiced amount across active transactions"),
            ("Total Amount Paid", float(kpis['total_paid']), CURRENCY_FORMAT, "Total payment amount recorded against billing"),
            ("Total STORY Amount", float(kpis['total_story']), CURRENCY_FORMAT, "Sum of recorded STORY values in source file"),
            ("STORY as % of Billing", float(kpis['story_percentage']) / 100.0, PERCENT_FORMAT, "Overall portfolio STORY margin ratio"),
            ("Total Billing Records", kpis['total_records'], NUMBER_FORMAT, "Total validated line-item transactions"),
            ("Unique Dealer Codes", kpis['unique_dealers'], NUMBER_FORMAT, "Distinct dealer accounts billed"),
            ("Unique Firms", kpis['unique_firms'], NUMBER_FORMAT, "Distinct firm names serviced"),
            ("Average Billing per Record", float(kpis['avg_billing_record']), CURRENCY_FORMAT, "Mean transaction value"),
            ("Average Billing per Dealer", float(kpis['avg_billing_dealer']), CURRENCY_FORMAT, "Mean dealer relationship value"),
            ("Highest Billing Month", kpis['highest_billing_month'], None, f"Total ₹{kpis['highest_billing_month_val']:,.2f}"),
            ("Highest Billing RM", kpis['top_billing_rm'], None, f"Total ₹{kpis['top_billing_rm_val']:,.2f}"),
            ("Highest STORY RM", kpis['top_story_rm'], None, f"Total ₹{kpis['top_story_rm_val']:,.2f}"),
            ("Latest MoM Billing Growth", float(kpis['mom_growth_pct']) / 100.0, PERCENT_FORMAT, "Month-on-Month billing momentum"),
            ("Reconciliation Difference (Billing - Paid)", float(kpis['payment_difference']), CURRENCY_FORMAT, "Unmatched or unpaid financial difference"),
            ("Data Quality Exceptions Flagged", kpis['quality_issues_count'], NUMBER_FORMAT, "Total automated audit anomalies identified")
        ]

        row_idx = 5
        for label, val, num_fmt, note in overview_metrics:
            ws_overview.append([label, val, note])
            c_label = ws_overview.cell(row=row_idx, column=1)
            c_val = ws_overview.cell(row=row_idx, column=2)
            c_note = ws_overview.cell(row=row_idx, column=3)

            c_label.font = BOLD_FONT
            c_val.font = DATA_FONT
            c_note.font = DATA_FONT
            for c in (c_label, c_val, c_note):
                c.border = THIN_BORDER
            if num_fmt:
                c_val.number_format = num_fmt
                if num_fmt in (CURRENCY_FORMAT, NUMBER_FORMAT, PERCENT_FORMAT):
                    c_val.alignment = Alignment(horizontal="right")
            row_idx += 1

        apply_auto_width(ws_overview)

        # -------------------------------------------------------------
        # Tab 2: RM Performance
        # -------------------------------------------------------------
        ws_rm = wb.create_sheet("RM Performance")
        ws_rm.append(["RELATIONSHIP MANAGER (RM) PERFORMANCE REPORT"])
        ws_rm.cell(row=1, column=1).font = TITLE_FONT
        ws_rm.append([f"Generated on {now_str}"])
        ws_rm.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_rm.append([])

        rm_headers = [
            "Relationship Manager", "Total Billing (₹)", "Total Paid (₹)", "Total STORY (₹)",
            "STORY %", "Transactions", "Unique Dealers", "Avg Billing/Tx (₹)",
            "Avg STORY/Tx (₹)", "% Share of Billing", "% Share of STORY"
        ]
        ws_rm.append(rm_headers)
        header_row = 4
        for col_idx in range(1, len(rm_headers) + 1):
            cell = ws_rm.cell(row=header_row, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        rm_data = AnalyticsService.get_rm_analytics(qs)
        curr_row = 5
        for r in rm_data:
            ws_rm.append([
                r['rm'],
                float(r['total_billing']),
                float(r['total_paid']),
                float(r['total_story']),
                float(r['story_percentage']) / 100.0,
                r['tx_count'],
                r['unique_dealers'],
                float(r['avg_billing']),
                float(r['avg_story_tx']),
                float(r['share_billing_pct']) / 100.0,
                float(r['share_story_pct']) / 100.0,
            ])
            for col_idx in range(1, len(rm_headers) + 1):
                c = ws_rm.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4, 8, 9):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (5, 10, 11):
                    c.number_format = PERCENT_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (6, 7):
                    c.number_format = NUMBER_FORMAT
                    c.alignment = Alignment(horizontal="right")
            curr_row += 1

        apply_auto_width(ws_rm)

        # -------------------------------------------------------------
        # Tab 3: Monthly Billing Trends
        # -------------------------------------------------------------
        ws_month = wb.create_sheet("Monthly Trends")
        ws_month.append(["MONTHLY BILLING & MOM GROWTH REPORT"])
        ws_month.cell(row=1, column=1).font = TITLE_FONT
        ws_month.append([f"Generated on {now_str}"])
        ws_month.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_month.append([])

        m_headers = [
            "Month", "Billing Amount (₹)", "Amount Paid (₹)", "STORY (₹)",
            "STORY %", "Transactions", "Unique Dealers", "Avg Billing (₹)",
            "Billing MoM Growth", "STORY MoM Growth"
        ]
        ws_month.append(m_headers)
        for col_idx in range(1, len(m_headers) + 1):
            cell = ws_month.cell(row=4, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        m_data = AnalyticsService.get_monthly_analytics(qs)
        curr_row = 5
        for m in m_data:
            ws_month.append([
                m['month_label'],
                float(m['total_billing']),
                float(m['total_paid']),
                float(m['total_story']),
                float(m['story_percentage']) / 100.0,
                m['tx_count'],
                m['unique_dealers'],
                float(m['avg_billing']),
                (float(m['billing_growth_pct']) / 100.0) if m['billing_growth_pct'] is not None else "-",
                (float(m['story_growth_pct']) / 100.0) if m['story_growth_pct'] is not None else "-",
            ])
            for col_idx in range(1, len(m_headers) + 1):
                c = ws_month.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4, 8):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx == 5:
                    c.number_format = PERCENT_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (6, 7):
                    c.number_format = NUMBER_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (9, 10):
                    if isinstance(c.value, (int, float)):
                        c.number_format = PERCENT_FORMAT
                        c.alignment = Alignment(horizontal="right")
                    else:
                        c.alignment = Alignment(horizontal="center")
            curr_row += 1

        apply_auto_width(ws_month)

        # -------------------------------------------------------------
        # Tab 4: STORY Analysis
        # -------------------------------------------------------------
        ws_story = wb.create_sheet("STORY Analysis")
        ws_story.append(["STORY MARGIN & ANOMALY ANALYSIS"])
        ws_story.cell(row=1, column=1).font = TITLE_FONT
        ws_story.append([f"Generated on {now_str}"])
        ws_story.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_story.append([])

        story_data = AnalyticsService.get_story_analytics(qs)
        ws_story.append(["STORY Metric", "Record Count", "Total STORY (₹)", "% of Population"])
        for col_idx in range(1, 5):
            c = ws_story.cell(row=4, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT

        total_tx = kpis['total_records']
        story_summary_rows = [
            ("Positive STORY (> ₹0)", story_data['pos_count'], float(story_data['pos_sum']), (story_data['pos_count'] / total_tx) if total_tx else 0),
            ("Zero STORY (= ₹0)", story_data['zero_count'], 0.0, (story_data['zero_count'] / total_tx) if total_tx else 0),
            ("Negative STORY (< ₹0)", story_data['neg_count'], float(story_data['neg_sum']), (story_data['neg_count'] / total_tx) if total_tx else 0),
        ]
        s_row = 5
        for lbl, cnt, amt, pct in story_summary_rows:
            ws_story.append([lbl, cnt, amt, pct])
            for col_idx in range(1, 5):
                c = ws_story.cell(row=s_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx == 2:
                    c.number_format = NUMBER_FORMAT
                elif col_idx == 3:
                    c.number_format = CURRENCY_FORMAT
                elif col_idx == 4:
                    c.number_format = PERCENT_FORMAT
            s_row += 1

        apply_auto_width(ws_story)

        # -------------------------------------------------------------
        # Tab 5: Category Breakdown
        # -------------------------------------------------------------
        ws_cat = wb.create_sheet("Category Analytics")
        ws_cat.append(["EXPERIENCE CATEGORY PERFORMANCE REPORT"])
        ws_cat.cell(row=1, column=1).font = TITLE_FONT
        ws_cat.append([f"Generated on {now_str}"])
        ws_cat.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_cat.append([])

        cat_headers = [
            "Category", "Billing Amount (₹)", "Amount Paid (₹)", "STORY (₹)",
            "STORY %", "Transactions", "Unique Dealers", "Share of Billing", "Share of STORY"
        ]
        ws_cat.append(cat_headers)
        for col_idx in range(1, len(cat_headers) + 1):
            c = ws_cat.cell(row=4, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center")

        cat_data = AnalyticsService.get_category_analytics(qs)
        curr_row = 5
        for cat in cat_data['category_list']:
            ws_cat.append([
                cat['standardized_category'],
                float(cat['billing']),
                float(cat['paid']),
                float(cat['story']),
                float(cat['story_pct']) / 100.0,
                cat['tx_count'],
                cat['unique_dealers'],
                float(cat['share_billing']) / 100.0,
                float(cat['share_story']) / 100.0,
            ])
            for col_idx in range(1, len(cat_headers) + 1):
                c = ws_cat.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (5, 8, 9):
                    c.number_format = PERCENT_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (6, 7):
                    c.number_format = NUMBER_FORMAT
                    c.alignment = Alignment(horizontal="right")
            curr_row += 1

        apply_auto_width(ws_cat)

        # -------------------------------------------------------------
        # Tab 6: Payment SPOC Analytics
        # -------------------------------------------------------------
        ws_spoc = wb.create_sheet("Payment SPOC")
        ws_spoc.append(["PAYMENT SPOC WORKLOAD & RECONCILIATION"])
        ws_spoc.cell(row=1, column=1).font = TITLE_FONT
        ws_spoc.append([f"Generated on {now_str}"])
        ws_spoc.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_spoc.append([])

        spoc_headers = ["Payment SPOC", "Billing Amount (₹)", "Amount Paid (₹)", "STORY (₹)", "Difference (₹)", "Transactions", "Unique Dealers", "STORY %"]
        ws_spoc.append(spoc_headers)
        for col_idx in range(1, len(spoc_headers) + 1):
            c = ws_spoc.cell(row=4, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT

        spoc_data = AnalyticsService.get_spoc_analytics(qs)
        curr_row = 5
        for sp in spoc_data:
            ws_spoc.append([
                sp['payment_spoc'],
                float(sp['billing']),
                float(sp['paid']),
                float(sp['story']),
                float(sp['diff']),
                sp['tx_count'],
                sp['unique_dealers'],
                float(sp['story_pct']) / 100.0,
            ])
            for col_idx in range(1, len(spoc_headers) + 1):
                c = ws_spoc.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4, 5):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (6, 7):
                    c.number_format = NUMBER_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx == 8:
                    c.number_format = PERCENT_FORMAT
                    c.alignment = Alignment(horizontal="right")
            curr_row += 1

        apply_auto_width(ws_spoc)

        # -------------------------------------------------------------
        # Tab 7: Transaction Detail (Filtered)
        # -------------------------------------------------------------
        ws_tx = wb.create_sheet("Detailed Transactions")
        tx_headers = [
            "Source Row", "Month", "Dealer Code", "Firm Name", "Experience Category",
            "RM", "Payment SPOC", "Billing Amount (₹)", "Amount Paid (₹)", "STORY (₹)", "Status"
        ]
        ws_tx.append(tx_headers)
        for col_idx in range(1, len(tx_headers) + 1):
            c = ws_tx.cell(row=1, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT

        # Limit to 5000 transactions for workbook performance
        curr_row = 2
        for tx in qs.order_by('id')[:5000]:
            ws_tx.append([
                tx.source_row_number,
                tx.month_label,
                tx.dealer_code,
                tx.firm_name,
                tx.standardized_category,
                tx.rm,
                tx.payment_spoc,
                float(tx.billing_amount),
                float(tx.amount_paid),
                float(tx.story_amount),
                tx.get_payment_status_display()
            ])
            for col_idx in range(1, len(tx_headers) + 1):
                c = ws_tx.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (8, 9, 10):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
            curr_row += 1

        apply_auto_width(ws_tx)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_pnl_workbook(filters=None):
        wb = openpyxl.Workbook()
        ws_monthly = wb.active
        ws_monthly.title = "Monthly P&L Statement"

        qs = AnalyticsService.get_filtered_queryset(filters)
        pnl_data = AnalyticsService.get_pnl_analytics(qs)
        now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

        # Tab 1: Monthly P&L
        ws_monthly.append(["HAVELLS BILLING - PROFIT & LOSS (P&L) STATEMENT"])
        ws_monthly.cell(row=1, column=1).font = TITLE_FONT
        ws_monthly.append([f"Report Generated: {now_str} | Active Scope: {pnl_data['tx_count']} Transactions"])
        ws_monthly.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_monthly.append([])

        m_headers = [
            "Month", "Gross Revenue (₹)", "Direct Fulfillment Cost (₹)", "Net Gross Profit / STORY (₹)",
            "Profit Margin %", "Cost Ratio %", "MoM Revenue Growth", "MoM Profit Growth",
            "Cumulative Profit (₹)", "Transactions", "Avg Profit / Tx (₹)"
        ]
        ws_monthly.append(m_headers)
        for col_idx in range(1, len(m_headers) + 1):
            cell = ws_monthly.cell(row=4, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        curr_row = 5
        for m in pnl_data['monthly_pnl']:
            ws_monthly.append([
                m['month_label'],
                float(m['revenue']),
                float(m['cost']),
                float(m['profit']),
                float(m['margin_pct']) / 100.0,
                float(m['cost_pct']) / 100.0,
                (float(m['rev_growth_pct']) / 100.0) if m['rev_growth_pct'] is not None else "-",
                (float(m['profit_growth_pct']) / 100.0) if m['profit_growth_pct'] is not None else "-",
                float(m['cumulative_profit']),
                m['tx_count'],
                float(m['avg_profit_tx'])
            ])
            for col_idx in range(1, len(m_headers) + 1):
                c = ws_monthly.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4, 9, 11):
                    c.number_format = CURRENCY_FORMAT
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (5, 6, 7, 8):
                    if isinstance(c.value, (int, float)):
                        c.number_format = PERCENT_FORMAT
                        c.alignment = Alignment(horizontal="right")
                    else:
                        c.alignment = Alignment(horizontal="center")
                elif col_idx == 10:
                    c.number_format = NUMBER_FORMAT
                    c.alignment = Alignment(horizontal="right")
            curr_row += 1

        # Total Row
        ws_monthly.append([
            "Grand Total",
            float(pnl_data['gross_revenue']),
            float(pnl_data['fulfillment_cost']),
            float(pnl_data['net_profit']),
            float(pnl_data['profit_margin_pct']) / 100.0,
            float(pnl_data['cost_ratio_pct']) / 100.0,
            "-", "-",
            float(pnl_data['net_profit']),
            pnl_data['tx_count'],
            float(pnl_data['net_profit'] / Decimal(str(pnl_data['tx_count']))) if pnl_data['tx_count'] else 0
        ])
        for col_idx in range(1, len(m_headers) + 1):
            c = ws_monthly.cell(row=curr_row, column=col_idx)
            c.font = BOLD_FONT
            c.fill = TOTAL_FILL
            c.border = THIN_BORDER
            if col_idx in (2, 3, 4, 9, 11):
                c.number_format = CURRENCY_FORMAT
                c.alignment = Alignment(horizontal="right")
            elif col_idx in (5, 6):
                c.number_format = PERCENT_FORMAT
                c.alignment = Alignment(horizontal="right")
            elif col_idx in (7, 8):
                c.alignment = Alignment(horizontal="center")
            elif col_idx == 10:
                c.number_format = NUMBER_FORMAT
                c.alignment = Alignment(horizontal="right")

        apply_auto_width(ws_monthly)

        # Tab 2: Yearly & Quarterly
        ws_yearly = wb.create_sheet("Annual & Quarterly P&L")
        ws_yearly.append(["ANNUAL AND QUARTERLY PROFIT & LOSS BREAKDOWN"])
        ws_yearly.cell(row=1, column=1).font = TITLE_FONT
        ws_yearly.append([f"Report Generated: {now_str}"])
        ws_yearly.cell(row=2, column=1).font = SUBTITLE_FONT
        ws_yearly.append([])

        y_headers = ["Financial Year / Quarter", "Revenue (₹)", "Fulfillment Cost (₹)", "Net Profit / Margin (₹)", "Profit Margin %"]
        ws_yearly.append(y_headers)
        for col_idx in range(1, len(y_headers) + 1):
            c = ws_yearly.cell(row=4, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT

        y_row = 5
        for y in pnl_data['yearly_pnl']:
            ws_yearly.append([f"Financial Year {y['financial_year']} (Full Year)", float(y['revenue']), float(y['cost']), float(y['profit']), float(y['margin_pct']) / 100.0])
            for col_idx in range(1, 6):
                c = ws_yearly.cell(row=y_row, column=col_idx)
                c.font = BOLD_FONT
                c.fill = TOTAL_FILL
                c.border = THIN_BORDER
                if col_idx in (2, 3, 4):
                    c.number_format = CURRENCY_FORMAT
                elif col_idx == 5:
                    c.number_format = PERCENT_FORMAT
            y_row += 1

            for q_name, q_obj in [("  - Q1 (Apr - Jun)", y['q1']), ("  - Q2 (Jul - Sep)", y['q2']), ("  - Q3 (Oct - Dec)", y['q3']), ("  - Q4 (Jan - Mar)", y['q4'])]:
                if q_obj['revenue'] > 0:
                    ws_yearly.append([q_name, float(q_obj['revenue']), float(q_obj['cost']), float(q_obj['profit']), float(q_obj['margin_pct']) / 100.0])
                    for col_idx in range(1, 6):
                        c = ws_yearly.cell(row=y_row, column=col_idx)
                        c.font = DATA_FONT
                        c.border = THIN_BORDER
                        if col_idx in (2, 3, 4):
                            c.number_format = CURRENCY_FORMAT
                        elif col_idx == 5:
                            c.number_format = PERCENT_FORMAT
                    y_row += 1

        apply_auto_width(ws_yearly)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
